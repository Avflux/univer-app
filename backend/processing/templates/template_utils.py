"""Utilitários para leitura e parsing de templates Excel e dados de resultados."""

import os
import openpyxl

from backend.core.logging.logger import get_logger
from backend.core.utils.paths import get_sheets_dir

log = get_logger(__name__)


def get_template_dir() -> str:
    """Retorna o diretório dos templates Excel (Sheets).

    Em modo frozen o diretório vem do bundle (``<_MEIPASS>/backend/Modules/
    Sheets``).
    """
    return get_sheets_dir()


def resolve_template(
    template_path: str | None,
    bay_id: str | None,
    bay_type: str | None = None,
    voltage_level: str | None = None,
) -> str | None:
    """Resolve o caminho do template. Prioridade: template_path explícito →
    bay_type+voltage_level → busca parcial pelo bay_id.
    """
    if template_path and os.path.isfile(template_path):
        log.info("_resolve_template: usando template_path explícito: %s", template_path)
        return template_path

    sheets_dir = get_template_dir()
    if not os.path.isdir(sheets_dir):
        log.warning("_resolve_template: diretório de templates não encontrado: %s", sheets_dir)
        return None

    available = [f for f in os.listdir(sheets_dir) if f.endswith(".xlsx")]
    log.info(
        "_resolve_template(bay=%s, type=%s, vt=%s): templates disponíveis=%s",
        bay_id, bay_type, voltage_level, available,
    )

    # Tenta por bay_type + voltage_level
    if bay_type and voltage_level:
        normalized_type = bay_type.upper().replace(" ", "")
        normalized_vt = voltage_level.upper().replace(" ", "")
        for fname in available:
            upper = fname.upper()
            if normalized_type in upper and normalized_vt in upper:
                log.info("_resolve_template: encontrado via type+vt: %s", fname)
                return os.path.join(sheets_dir, fname)

    # Fallback por bay_id
    if bay_id:
        for fname in available:
            if bay_id.lower() in fname.lower():
                log.info("_resolve_template: encontrado via bay_id: %s", fname)
                return os.path.join(sheets_dir, fname)

    log.warning(
        "_resolve_template: nenhum template encontrado para bay=%s type=%s vt=%s",
        bay_id, bay_type, voltage_level,
    )
    return None


def get_template_connections(template_path, measurement_type="Costura", seccionadora_count=0, sincronizador_enabled=True):
    """
    Escaneia o template e retorna todas as conexões (origem, destino) que precisam de distância.
    Exclui conexões com distância fixa no template.
    """
    if not template_path or not os.path.exists(template_path):
        return set()

    template_wb = openpyxl.load_workbook(template_path)
    connections = set()

    for sheet_name in template_wb.sheetnames:
        if sheet_name == "SECCIONADORA":
            continue
        if sheet_name == "SINCRONIZADOR" and not sincronizador_enabled:
            continue
        if sheet_name == "COSTURA" and measurement_type == "Caixa":
            continue
        if sheet_name == "CAIXA CA" and measurement_type == "Costura":
            continue
        # Em "Sem Alim." não plotamos nem COSTURA nem CAIXA CA — o xlsm
        # final fica sem essas duas abas independentemente de qualquer
        # estado salvo (checkbox, alias, etc.).
        if measurement_type == "Sem Alim." and sheet_name in ("COSTURA", "CAIXA CA"):
            continue

        template_sheet = template_wb[sheet_name]

        for row in range(2, template_sheet.max_row + 1):
            cell_d = template_sheet[f'D{row}'].value
            cell_e = template_sheet[f'E{row}'].value

            if not cell_d or not cell_e:
                continue

            if sheet_name == "CAIXA CA" and cell_d != "CAIXA DE DISTRIBUIÇÃO CA":
                continue

            if measurement_type == "Caixa":
                g_value = template_sheet[f'G{row}'].value
                if g_value and str(g_value).strip() == '*':
                    continue

            template_distance = template_sheet[f'F{row}'].value
            if template_distance is not None and isinstance(template_distance, (int, float)) and template_distance > 0:
                continue

            if cell_e == "SECCIONADORA":
                for i in range(1, seccionadora_count + 1):
                    connections.add((cell_d, f"SECCIONADORA-{i}"))
            elif cell_d == "SECCIONADORA":
                for i in range(1, seccionadora_count + 1):
                    connections.add((f"SECCIONADORA-{i}", cell_e))
            else:
                connections.add((cell_d, cell_e))

    if seccionadora_count > 0 and "SECCIONADORA" in template_wb.sheetnames:
        template_secc = template_wb["SECCIONADORA"]
        for row in range(2, template_secc.max_row + 1):
            cell_d = template_secc[f'D{row}'].value
            cell_e = template_secc[f'E{row}'].value
            if not cell_d or not cell_e:
                continue
            template_distance = template_secc[f'F{row}'].value
            if template_distance is not None and isinstance(template_distance, (int, float)) and template_distance > 0:
                continue
            for i in range(1, seccionadora_count):
                d = cell_d.replace("SECCIONADORA", f"SECCIONADORA-{i}") if "SECCIONADORA" in cell_d else cell_d
                e = cell_e.replace("SECCIONADORA", f"SECCIONADORA-{i+1}") if "SECCIONADORA" in cell_e else cell_e
                connections.add((d, e))

    template_wb.close()
    return connections
