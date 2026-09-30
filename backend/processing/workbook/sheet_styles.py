"""Utilitários de estilização centralizados para planilhas Excel."""

from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.styles.numbers import FORMAT_NUMBER_00


HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
TITLE_FILL = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
CENTER_ALIGNMENT = Alignment(horizontal='center', vertical='center')


def apply_styles(sheet, columns=None):
    """
    Aplica estilos na planilha: ajusta largura de colunas, formata header
    e formata números como 2 casas decimais na coluna F (por padrão).

    Args:
        sheet: planilha openpyxl.
        columns: lista de letras de colunas a formatar. Padrão: A-F.
    """
    columns = columns or ['A', 'B', 'C', 'D', 'E', 'F']

    for col in columns:
        max_length = 0
        column = sheet[col]
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        sheet.column_dimensions[col].width = max_length + 2

    for col in columns:
        header_cell = sheet[f'{col}1']
        header_cell.font = HEADER_FONT
        header_cell.fill = HEADER_FILL
        header_cell.border = THIN_BORDER
        header_cell.alignment = CENTER_ALIGNMENT

    if 'F' in columns:
        for row in range(2, sheet.max_row + 1):
            cell = sheet[f'F{row}'].value
            if cell is not None:
                try:
                    float(str(cell).replace(',', '.'))
                    sheet[f'F{row}'].number_format = FORMAT_NUMBER_00
                except (ValueError, TypeError):
                    pass


def apply_total_sheet_styles(sheet, columns=None):
    """
    Versão específica para abas de totais (2 colunas A-B), com formato de número na coluna B.
    """
    columns = columns or ['A', 'B']
    apply_styles(sheet, columns)

    for row in range(2, sheet.max_row + 1):
        cell = sheet[f'B{row}'].value
        if cell is not None:
            try:
                float(str(cell).replace(',', '.'))
                sheet[f'B{row}'].number_format = FORMAT_NUMBER_00
            except (ValueError, TypeError):
                pass
