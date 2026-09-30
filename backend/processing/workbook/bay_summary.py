"""Resumo por bay (uma seção TOTAL)."""

from typing import Callable, Dict, List
from openpyxl.workbook.workbook import Workbook

SUMMARY_COLUMNS_FOR_TOTAL = ('A', 'B')

# Implementações vêm das injeções para isolar o ciclo de fontes/bordas
# (HEADER_FONT, HEADER_FILL, THIN_BORDER, CENTER_ALIGNMENT).


def build_summary_sheet(
    workbook: Workbook,
    aggregate_cable_totals_fn: Callable[[Workbook], Dict[str, float]],
    styles: dict,
    title: str = "TOTAL",
) -> None:
    """Cria a aba ``title`` agregando totais por formação.

    ``styles`` deve prover: HEADER_FONT, HEADER_FILL, THIN_BORDER, CENTER_ALIGNMENT.
    """
    sheet = workbook.create_sheet(title=title)
    sheet['A1'] = "Formação"
    sheet['B1'] = "Distância Total (m)"
    for col in SUMMARY_COLUMNS_FOR_TOTAL:
        header_cell = sheet[f'{col}1']
        header_cell.font = styles['HEADER_FONT']
        header_cell.fill = styles['HEADER_FILL']
        header_cell.border = styles['THIN_BORDER']
        header_cell.alignment = styles['CENTER_ALIGNMENT']

    cable_totals = aggregate_cable_totals_fn(workbook)
    row = 2
    for cable_type, total_distance in sorted(cable_totals.items()):
        sheet[f'A{row}'].value = cable_type
        sheet[f'B{row}'].value = total_distance
        for col in SUMMARY_COLUMNS_FOR_TOTAL:
            cell = sheet[f'{col}{row}']
            cell.border = styles['THIN_BORDER']
            cell.alignment = styles['CENTER_ALIGNMENT']
        row += 1


def build_multi_bays_summary(
    workbook: Workbook,
    bay_data: List[dict],
    headers: List[str],
    styles: dict,
    get_index: int = 0,
    title: str = "SUMÁRIO GERAL",
) -> None:
    """Cria uma aba ``title`` listando uma linha por bay."""
    sheet = workbook.create_sheet(title=title, index=get_index)
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col)
        cell.value = header
        cell.font = styles['HEADER_FONT']
        cell.fill = styles['HEADER_FILL']
        cell.border = styles['THIN_BORDER']
        cell.alignment = styles['CENTER_ALIGNMENT']

    for row, bay in enumerate(bay_data, 2):
        for col, accessor in enumerate(ACCESSORS, 1):
            sheet.cell(row=row, column=col).value = accessor(bay)
            sheet.cell(row=row, column=col).border = styles['THIN_BORDER']
            sheet.cell(row=row, column=col).alignment = styles['CENTER_ALIGNMENT']

    for col_idx in range(1, len(headers) + 1):
        from openpyxl.utils import get_column_letter
        column = get_column_letter(col_idx)
        max_length = 0
        for cell in sheet[column]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        sheet.column_dimensions[column].width = max_length + 2


ACCESSORS = [
    lambda bay: bay['name'],
    lambda bay: bay['type'],
    lambda bay: bay['voltage'],
    lambda bay: bay['measurement_type'],
    lambda bay: len(bay['results']),
    lambda bay: bay['seccionadora_count'],
    lambda bay: bay.get('tc_phases', 1),
    lambda bay: bay.get('tp_phases', 1),
    lambda bay: "Sim" if bay.get('merging_unit_enabled', False) else "Não",
]
