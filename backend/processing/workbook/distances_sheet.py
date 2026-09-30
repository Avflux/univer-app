"""Geração da planilha Excel da aba "Distâncias" de um bay.

Separa a construção do workbook (estilos, cabeçalhos, linhas) da
persistência em disco, que fica a cargo do chamador (sidecar).
"""

from __future__ import annotations

import openpyxl
from openpyxl.styles.numbers import FORMAT_NUMBER_00


def coerce_distance(value: object) -> float | str:
    """Normaliza a distância recebida para um valor numérico quando possível.

    A célula precisa ser **numérica** para o Excel entrar nas somas/fórmulas.
    O separador decimal exibido (vírgula no Excel em pt-BR, ponto em en-US) é
    renderizado pelo próprio Excel a partir do ``number_format`` — gravar o
    texto ``"12,50"`` mostraria a vírgula, mas viraria texto e não calcularia.

    Aceita ``12.5``, ``"12.5"`` e ``"12,5"`` (distâncias salvas no .md podem
    usar vírgula). Texto não numérico é preservado como está.

    Args:
        value: Valor bruto da distância.

    Returns:
        ``float`` quando numérico, ``""`` quando vazio, ou o texto original.
    """
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return ""
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return text


def build_distances_workbook(distances: list[dict]) -> openpyxl.Workbook:
    """Monta o workbook da aba "Distâncias".

    Recebe o array de distâncias serializado pelo frontend:
      [{ "id": "dist-1", "origem": "PAINEL", "destino": "DISJUNTOR",
         "distancia": "12.5" }, ...]

    A ordem das linhas no arquivo é igual à ordem recebida (já ordenada
    pelo frontend conforme a ordem de seleção dos pontos).

    Args:
        distances: Distâncias no formato [{id, origem, destino, distancia}].

    Returns:
        Workbook aberto, pronto para ``save`` pelo chamador.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Distâncias"

    # Cabeçalhos
    headers = ["Origem", "Destino", "Distância (m)"]
    header_font = openpyxl.styles.Font(bold=True, size=11)
    header_fill = openpyxl.styles.PatternFill(
        start_color="E0E0E0", end_color="E0E0E0", fill_type="solid",
    )
    border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style="thin"),
        right=openpyxl.styles.Side(style="thin"),
        top=openpyxl.styles.Side(style="thin"),
        bottom=openpyxl.styles.Side(style="thin"),
    )
    center = openpyxl.styles.Alignment(
        horizontal="center", vertical="center",
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    # Linhas de dados — mantém a ordem recebida do frontend
    for row_idx, dist in enumerate(distances, 2):
        origem = dist.get("origem", "")
        destino = dist.get("destino", "")
        distancia = coerce_distance(dist.get("distancia"))

        ws.cell(row=row_idx, column=1, value=origem).border = border
        ws.cell(row=row_idx, column=2, value=destino).border = border
        distance_cell = ws.cell(row=row_idx, column=3, value=distancia)
        distance_cell.border = border
        distance_cell.alignment = center
        if isinstance(distancia, float):
            # Duas casas decimais: o Excel mostra "12,50" em pt-BR.
            distance_cell.number_format = FORMAT_NUMBER_00

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 18

    return wb
