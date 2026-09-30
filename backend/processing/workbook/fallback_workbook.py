"""Workbook de fallback quando não há template disponível."""

from typing import List

import openpyxl
from openpyxl.styles.numbers import FORMAT_NUMBER_00
from openpyxl.workbook.workbook import Workbook

from backend.processing.constants import SHEET_DISTANCIAS, Connection


def create_new_workbook(results_data: List[Connection]) -> Workbook:
    """Cria um workbook simples apenas com as distâncias (sem abas de template).

    A distância vai como número formatado em duas casas decimais, para o
    Excel calcular e exibir com o separador decimal do idioma ("12,50" em
    pt-BR).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_DISTANCIAS
    ws['D1'] = "Origem"
    ws['E1'] = "Destino"
    ws['F1'] = "Distância (m)"
    for idx, (origin, dest, dist) in enumerate(results_data, start=2):
        ws[f'D{idx}'] = origin
        ws[f'E{idx}'] = dest
        ws[f'F{idx}'] = dist
        if isinstance(dist, (int, float)):
            ws[f'F{idx}'].number_format = FORMAT_NUMBER_00
    return wb
