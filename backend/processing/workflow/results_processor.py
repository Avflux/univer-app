"""Orquestrador do fluxo de geração do workbook de um bay (process_results)."""

import os
import re
from typing import Any

import openpyxl
from openpyxl.workbook.workbook import Workbook

from backend.core.logging.logger import get_logger
from backend.processing.constants import (
    CAIXA_DISTRIBUICAO_CA,
    DISJUNTOR,
    MEASUREMENT_CAIXA,
    MEASUREMENT_COSTURA,
    SECCIONADORA,
    TC,
    TOTAL,
    TP,
    Connection,
)
from backend.processing.workbook.bay_summary import build_summary_sheet
from backend.processing.workbook.cable_totals import aggregate_cable_totals
from backend.processing.workbook.distance_manager import DistanceManager
from backend.processing.workbook.fallback_workbook import create_new_workbook
from backend.processing.workbook.sheet_styles import apply_total_sheet_styles
from backend.processing.workflow.measurement_rules import (
    selected_equipment,
    should_process_sheet,
)
from backend.processing.workflow.seccionadora_processor import (
    has_other_seccionadora_connections,
    process_seccionadora,
)
from backend.processing.workflow.sheet_processor import process_sheet
from backend.processing.workflow.sheet_routing import special_sheet_kind
from backend.processing.workflow.tc_tp_processor import process_tc_tp

log = get_logger(__name__)


def create_summary_sheet(processor, workbook: Workbook) -> None:
    """Cria a aba TOTAL agregando totais por formação."""
    build_summary_sheet(
        workbook,
        aggregate_cable_totals_fn=aggregate_cable_totals,
        styles=processor._summary_styles(),
        title=TOTAL,
    )
    apply_total_sheet_styles(workbook[TOTAL])


def process_results(
    processor,
    results_data: list[Connection],
    seccionadora_count: int = 0,
    measurement_type: str = MEASUREMENT_COSTURA,
    template_path: str | None = None,
    sincronizador_enabled: bool = True,
    tc_phases: int = 1,
    tp_phases: int = 1,
    merging_unit_enabled: bool = False,
    bay_name: str | None = None,
    point_aliases: dict[str, str] | None = None,
    cable_data: list[dict[str, Any]] | None = None,
) -> Workbook:
    """Gera o workbook consolidado de um bay a partir do template e das conexões.

    ``processor`` deve expor: cable_resolver, cable_calculator, distance_manager
    e ``_summary_styles`` (ExcelProcessor).
    """
    processor.cable_resolver.clear_cache()
    if cable_data and bay_name:
        processor.cable_resolver.load_bay_cable_data(bay_name, cable_data)
    processor.distance_manager = DistanceManager(point_aliases or {})

    # Repassa os aliases para a calculadora de cabos, indexados por bay.
    if bay_name:
        processor.cable_calculator.bay_aliases[bay_name] = (
            processor.distance_manager.point_aliases
        )
    log.debug(
        "Iniciando process_results: measurement_type=%s, seccionadora_count=%s",
        measurement_type,
        seccionadora_count,
    )
    log.debug(
        "TC phases: %s, TP phases: %s, MERGING UNIT: %s",
        tc_phases,
        tp_phases,
        "ativado" if merging_unit_enabled else "desativado",
    )
    log.debug("results_data contém %s conexões", len(results_data))

    processor.cable_resolver.sync_cache_with_calculator()

    if not template_path or not os.path.exists(template_path):
        log.warning("Template file not found: %s", template_path)
        return create_new_workbook(results_data)

    template_wb = openpyxl.load_workbook(template_path)
    new_wb = openpyxl.Workbook()
    new_wb.remove(new_wb.active)

    selected = selected_equipment(results_data)

    # Em "Sem Alim." não há aba CAIXA CA, então CAIXA DE DISTRIBUIÇÃO
    # CA também não deve entrar como equipamento selecionado.
    if measurement_type == MEASUREMENT_CAIXA:
        selected.add(CAIXA_DISTRIBUICAO_CA)

    # MEASUREMENT_NONE: nem COSTURA nem CAIXA CA entram — apenas os
    # equipamentos explicitamente selecionados via checkbox.

    if seccionadora_count > 0 and SECCIONADORA in template_wb.sheetnames:
        template_secc_sheet = template_wb[SECCIONADORA]

        for i in range(1, seccionadora_count + 1):
            new_sheet = new_wb.create_sheet(title=f"{SECCIONADORA}-{i}")

            if i == seccionadora_count and not has_other_seccionadora_connections(
                template_secc_sheet
            ):
                new_wb.remove(new_sheet)
                continue

            if not process_seccionadora(
                processor,
                template_secc_sheet,
                new_sheet,
                results_data,
                i,
                measurement_type,
                merging_unit_enabled,
                bay_name,
            ):
                new_wb.remove(new_sheet)

    # --- Coleta os pares de SECCIONADORA-N para excluir do DISJUNTOR ---
    # Qualquer conexão cujo par (frozenset) envolva um endpoint SECCIONADORA-N
    # já foi (ou será) escrita nas abas SECCIONADORA-N.  Ao processar a aba
    # DISJUNTOR, esses pares devem ser ignorados para evitar duplicidade.
    _secc_re = re.compile(r"SECCIONADORA-\d+", re.IGNORECASE)
    secc_pairs: set[frozenset[str]] = set()
    if seccionadora_count > 0:
        for orig, dest, _ in results_data:
            if _secc_re.search(orig) or _secc_re.search(dest):
                secc_pairs.add(frozenset((orig, dest)))

    for sheet_name in template_wb.sheetnames:
        if sheet_name == SECCIONADORA:
            continue
        if not should_process_sheet(
            sheet_name, measurement_type, selected, sincronizador_enabled
        ):
            continue
        # Filtro por modo de alimentação:
        # - "Sem Alim." não plota nem COSTURA nem CAIXA CA.
        # - "Caixa" pula COSTURA (e mantém CAIXA CA).
        # - "Costura" pula CAIXA CA (e mantém COSTURA).
        template_sheet = template_wb[sheet_name]
        new_sheet = new_wb.create_sheet(title=sheet_name)

        kind = special_sheet_kind(sheet_name, selected)

        special = {
            "tc": (TC, tc_phases),
            "tp": (TP, tp_phases),
        }

        if kind in special:
            sheet_type, phases = special[kind]

            if process_tc_tp(
                processor,
                template_sheet,
                new_sheet,
                results_data,
                sheet_type,
                phases,
                merging_unit_enabled,
                bay_name,
            ):
                continue

        # Para a aba DISJUNTOR: remove da lista de conexões os pares que
        # já foram registrados nas tabelas SECCIONADORA-N (via única — o
        # cabo é lançado em apenas uma ponta, não deve ser duplicado).
        effective_results = (
            [
                conn
                for conn in results_data
                if frozenset((conn[0], conn[1])) not in secc_pairs
            ]
            if sheet_name == DISJUNTOR and secc_pairs
            else results_data
        )

        if not process_sheet(
            processor,
            template_sheet,
            new_sheet,
            effective_results,
            None,
            measurement_type,
            merging_unit_enabled,
            bay_name,
        ):
            new_wb.remove(new_sheet)

    create_summary_sheet(processor, new_wb)
    return new_wb
