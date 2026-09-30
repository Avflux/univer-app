"""Finalização do workbook multi-bay: SUMÁRIO GERAL → abas por bay → TOTAL GERAL."""

from typing import Callable, Dict, List

import openpyxl
from openpyxl.workbook.workbook import Workbook

from backend.processing.constants import (
    BayData,
    DATA_COLUMNS,
    SUMARIO_GERAL,
    TOTAL,
    TOTAL_GERAL,
)
from backend.processing.workbook.bay_section_copier import assemble_bay_sheet
from backend.processing.workbook.bay_summary import build_multi_bays_summary
from backend.processing.workbook.total_geral_formatter import (
    auto_fit_columns,
    is_shielded,
    standardize_formation,
    write_adjusted_section,
    write_cable_section,
)
from backend.processing.workbook.sheet_styles import CENTER_ALIGNMENT, HEADER_FILL, HEADER_FONT, THIN_BORDER, TITLE_FILL


SUMMARY_HEADERS = [
    "Bay", "Tipo", "Tensão", "Tipo de Medição", "Total de Conexões",
    "Seccionadoras", "Enrolamentos TC", "Enrolamentos TP", "Merging Unit",
]


def multi_bay_styles() -> dict:
    """Retorna estilos compartilhados pela montagem do workbook multi-bay."""
    return {
        'HEADER_FONT': HEADER_FONT, 'HEADER_FILL': HEADER_FILL,
        'TITLE_FILL': TITLE_FILL, 'THIN_BORDER': THIN_BORDER,
        'CENTER_ALIGNMENT': CENTER_ALIGNMENT,
    }


def prepare_multi_bay_workbook(bay_data: List[BayData]):
    """Cria workbook inicial, estilos e callback de montagem por bay."""
    styles = multi_bay_styles()
    return build_summary_workbook(bay_data, styles), styles, assemble_default_per_bay(styles)


def assemble_default_per_bay(styles: dict) -> Callable[[Workbook, BayData], None]:
    """Callback padrão de montagem por bay — usa bay['_workbook'] e bay['totals']."""
    def _callback(final_wb: Workbook, bay: BayData) -> None:
        bay_wb = bay['_workbook']
        assemble_bay_sheet(
            final_wb, bay_wb, bay,
            total_label=TOTAL, styles=styles, data_columns=DATA_COLUMNS,
        )
    return _callback


def build_summary_workbook(bay_data: List[BayData], styles: dict) -> Workbook:
    """Cria workbook vazio com SUMÁRIO GERAL já preenchido."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_multi_bays_summary(
        wb, bay_data, SUMMARY_HEADERS, styles,
        get_index=0, title=SUMARIO_GERAL,
    )
    return wb


def append_total_geral_sheet(workbook: Workbook, bay_data: List[BayData],
                             styles: dict, logger=None) -> None:
    """Adiciona a aba TOTAL_GERAL ao workbook final."""
    from openpyxl.styles import Font

    if logger is not None:
        logger.debug("==== Criando aba de Total Geral ====")

    sheet = workbook.create_sheet(title=TOTAL_GERAL, index=1)

    sheet['A1'] = "TOTAL GERAL DE TODOS OS BAYS"
    sheet.merge_cells('A1:C1')
    header_cell = sheet['A1']
    header_cell.font = Font(bold=True, size=14)
    header_cell.fill = styles['TITLE_FILL']
    header_cell.alignment = styles['CENTER_ALIGNMENT']
    for col in ('A', 'B', 'C'):
        sheet[f'{col}1'].border = styles['THIN_BORDER']

    sheet['E1'] = "TOTAL GERAL AJUSTADO (10%)"
    sheet.merge_cells('E1:F1')
    ajustado_header = sheet['E1']
    ajustado_header.font = Font(bold=True, size=14)
    ajustado_header.fill = styles['TITLE_FILL']
    ajustado_header.alignment = styles['CENTER_ALIGNMENT']
    for col in ('E', 'F'):
        sheet[f'{col}1'].border = styles['THIN_BORDER']

    shielded_cables: Dict[str, Dict] = {}
    unshielded_cables: Dict[str, Dict] = {}
    for bay in bay_data:
        bay_name = bay['name']
        if 'totals' not in bay:
            if logger is not None:
                logger.warning("Aviso: Bay %s não tem totais calculados", bay_name)
            continue
        bay_totals = bay['totals']
        for cable_type, total_distance in bay_totals.items():
            if not cable_type or total_distance == 0:
                continue
            target_dict = shielded_cables if is_shielded(cable_type) else unshielded_cables
            normalized_type = standardize_formation(cable_type)
            if normalized_type not in target_dict:
                target_dict[normalized_type] = {
                    'total': 0.0,
                    'bays': set(),
                    'original_types': set(),
                }
            target_dict[normalized_type]['total'] += total_distance
            target_dict[normalized_type]['bays'].add(bay_name)
            target_dict[normalized_type]['original_types'].add(cable_type)

    current_row = 3
    if shielded_cables:
        current_row = write_cable_section(
            sheet, "CABOS BLINDADOS", shielded_cables, current_row,
            styles['TITLE_FILL'], styles['HEADER_FILL'], styles['THIN_BORDER'],
            styles['CENTER_ALIGNMENT'],
        )
    if unshielded_cables:
        current_row = write_cable_section(
            sheet, "CABOS NÃO BLINDADOS", unshielded_cables, current_row,
            styles['TITLE_FILL'], styles['HEADER_FILL'], styles['THIN_BORDER'],
            styles['CENTER_ALIGNMENT'],
        )

    current_row = 3
    if shielded_cables:
        current_row = write_adjusted_section(
            sheet, "CABOS BLINDADOS", shielded_cables, current_row,
            styles['TITLE_FILL'], styles['HEADER_FILL'], styles['THIN_BORDER'],
            styles['CENTER_ALIGNMENT'],
        )
    if unshielded_cables:
        current_row += 1
    if unshielded_cables:
        current_row = write_adjusted_section(
            sheet, "CABOS NÃO BLINDADOS", unshielded_cables, current_row,
            styles['TITLE_FILL'], styles['HEADER_FILL'], styles['THIN_BORDER'],
            styles['CENTER_ALIGNMENT'],
        )

    auto_fit_columns(sheet, (1, 2, 3))
    auto_fit_columns(sheet, (5, 6))
